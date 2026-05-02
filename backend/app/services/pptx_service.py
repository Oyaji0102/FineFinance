from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
import io

def _add_title_and_text(slide, title_text, bullet_points):
    title = slide.shapes.title
    title.text = title_text
    body = slide.placeholders[1]
    tf = body.text_frame
    for i, pt in enumerate(bullet_points):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = pt
        p.level = 0
        p.font.size = Pt(18)

def create_firm_presentation(firm_data: dict, ai_analysis: dict, ai_summary: str = "Özet bulunamadı.") -> bytes:
    prs = Presentation()
    primary_color = RGBColor(0, 51, 102) 
    
    # 1. KAPAK SLAYTI
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    title1 = slide1.shapes.title
    title1.text = f"{firm_data.get('name', 'Firma Adı Yok')}\nFinansal Analiz ve Ön Sunum"
    title1.text_frame.paragraphs[0].font.color.rgb = primary_color
    title1.text_frame.paragraphs[0].font.bold = True
    slide1.placeholders[1].text = "FineFinance AI Tarafından Üretilmiştir"
    
    # 2. AI KISA YORUM / ÖZET SLAYT (İstenen 5. Madde)
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    # ai_summary düz metin olduğu için tek elemanlı bir liste olarak veriyoruz
    _add_title_and_text(slide2, "1. AI Finansal Yönetici Özeti", [ai_summary])
    
    # 3. TEMEL MALİ VERİLER (İstenen 1. Madde)
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])
    slide3.shapes.title.text = "2. Mali Yapı Özeti ve Temel Veriler"
    
    rows, cols = 5, 2
    table_shape = slide3.shapes.add_table(rows, cols, Inches(1), Inches(1.5), Inches(8), Inches(2.5))
    table = table_shape.table
    table.columns[0].width = Inches(4)
    table.columns[1].width = Inches(4)
    
    data = [
        ("Finansal Metrik", "Güncel Değer"),
        ("Tahmini Yıllık Ciro", f"{firm_data.get('revenue', '0')} TL"),
        ("Toplam Varlıklar (Aktif)", f"{firm_data.get('assets', '0')} TL"),
        ("Kısa Vadeli Borçlar", f"{firm_data.get('short_term_debt', '0')} TL"),
        ("Özkaynaklar", f"{firm_data.get('equity', '0')} TL")
    ]
    
    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(val)
            if r_idx == 0:
                cell.text_frame.paragraphs[0].font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(220, 220, 220)

    # 4. FİNANSAL RASYOLAR (İstenen 3. Madde)
    slide4 = prs.slides.add_slide(prs.slide_layouts[1])
    ratios = [
        f"Likidite (Cari Oran): {firm_data.get('ratios', {}).get('current_ratio', 'Belirtilmedi')}",
        f"Karlılık (Net Kar Marjı): {firm_data.get('ratios', {}).get('net_profit_margin', 'Belirtilmedi')}",
        f"Kaldıraç (Borç/Özkaynak Oranı): {firm_data.get('ratios', {}).get('debt_to_equity', 'Belirtilmedi')}"
    ]
    _add_title_and_text(slide4, "3. Finansal Rasyolar", ratios)
    
    # 5. BANKADAKİ GENEL DURUM VE LİMİTLER (İstenen 2. Madde)
    slide5 = prs.slides.add_slide(prs.slide_layouts[1])
    banks = firm_data.get("banks", ["Banka ve limit verisi bulunamadı."])
    _add_title_and_text(slide5, "4. Banka Durumu ve Kredi Limitleri", banks)
    
    # 6. TRENDLER GRAFİĞİ (İstenen 4. Madde)
    slide6 = prs.slides.add_slide(prs.slide_layouts[5])
    slide6.shapes.title.text = "5. Finansal Trendler (Geçmiş Yıllar Büyüme)"
    
    chart_data = CategoryChartData()
    chart_data.categories = ['2021', '2022', '2023', 'Güncel Yıl']
    trends = firm_data.get("trends", [0, 0, 0, 0])
    chart_data.add_series('Ciro (Milyon TL)', tuple(trends))
    
    # Line chart (Çizgi grafiği) kullanarak büyüme trendini gösteriyoruz
    slide6.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(1), Inches(2), Inches(8), Inches(4.5), chart_data)

    pptx_stream = io.BytesIO()
    prs.save(pptx_stream)
    pptx_stream.seek(0)
    return pptx_stream.read()
